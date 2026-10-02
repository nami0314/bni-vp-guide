from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs", "launch_team_management_v5.xlsx")
F = "Meiryo"
NAVY = "1F3A5F"; BLUE = "2E75B6"; LBLUE = "DDEBF7"; YEL = "FFF2CC"; GRAY = "F2F2F2"
thin = Side(style="thin", color="BFBFBF"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
def font(**k): return Font(name=F, **k)
def fill(c): return PatternFill("solid", fgColor=c)
def hdr(c, color=NAVY):
    c.font = font(bold=True, color="FFFFFF", size=10); c.fill = fill(color)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BD
    return c
def cell(ws, r, c, v=None, inp=False, fmt=None, bold=False, wrap=False, center=False):
    x = ws.cell(r, c)
    if v is not None: x.value = v
    x.font = font(size=10, bold=bold); x.border = BD
    x.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal="center" if center else None)
    if inp: x.fill = fill(YEL)
    if fmt: x.number_format = fmt
    return x
def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = font(bold=True, size=14, color=NAVY)
    if sub: ws["A2"] = sub; ws["A2"].font = font(size=9, color="595959")
def widths(ws, ws_list):
    for i, w in enumerate(ws_list, 1): ws.column_dimensions[L(i)].width = w

LDS = [("竹尾純子","JJ"),("福重孝一","Believe"),("矢野貴之","チェキン"),("榎田秀一","JJ"),
       ("世良真貴男","Believe"),("松永健志","チェキン"),("中村剛","JJ")]
names = [n for n,_ in LDS]
wb = Workbook()

# ================= 使い方 =================
ws = wb.active; ws.title = "使い方"
title(ws, "ローンチチーム管理表｜使い方と運用ルール", "黄色セル＝入力欄／白セル＝自動計算（触らない）　※メニュー「🚀ローンチ管理」はGAS（Code.gs）を入れると表示されます")
rows = [
("■ 全体の流れ",""),
("","声かけ → 【CCS追っかけ】で入会まで追う → 申込書・入金・同意書の3つに✓ → 自動で【入会後行動リスト】に本人とDNAのタスクが作られる → 毎日決まった時刻に、遅れている人・LINE未参加の人をメールでアラート"),
("■ シート構成",""),
("📊ダッシュボード","全ローンチの進捗・信号・体制・入会後タスクの遅れを一覧（自動計算）。週次ローンチMTGはこの画面で進行。"),
("全体管理","従来の全体管理表と同じ並び。各LDシートから自動集約（入力不要）。"),
("週次推移","10/5スタートの13週（年末まで）。声かけ・面談・CCS実施・正式化の週ごとの推移。開始日は設定シートで変更可。"),
("CCS追っかけ","リストアップ〜入会までの台帳。1人1行。担当（リストアップしたメンバー）がCCSまで追いかける。現状・次のアクション・期日を入れ、申込書・振込控え・プリコア同意書の3点✓で『✅正式』。"),
("担当別CCS","ローンチ×担当メンバーごとの、リストアップ数・お誘い中・CCS済・入会数（GASが毎日更新）。"),
("入会後行動リスト","入会者ごとの行動（本人のタスク＋DNAのタスク）。GASが自動で作成。完了したら✓。LINEグループに入ったのを確認したら『LINEグループに参加』に✓。"),
("入会後進捗","入会者×行動のマトリクス（GASが毎日自動更新）。誰が何で止まっているかを一目で確認。"),
("日次ログ","毎日の活動を1ローンチ1行で記録。"),
("各LDシート","LDが自分のチーム形成・フェーズ情報を入力。"),
("行動マスタ","入会後の流れ（ハルモニアのメンバー表の項目）。本人/DNA・担当ロール・期限=正式化日+何日・本人リマインド要否。ここを変えると次の入会者から反映。"),
("通知先","ローンチごとのLDメール・CC・メンバー用LINEグループ招待URLと、DNAのメールアドレス。"),
("設定","期限日数・アラートの判定日数・プルダウン項目。"),
("アラートログ","メール送信履歴（自動）。"),
("■ 毎日（担当メンバー・LD・5分）",""),
("","① 担当メンバー：届いたメールを見て見込み客に連絡 → CCS追っかけの現状・次のアクション・期日を更新（編集すると最終接触日が自動で今日になります）"),
("","② 入会後行動リストで終わったタスクに✓（完了日は自動）　③ 日次ログに1行（0の日も0で入力）"),
("■ 毎週（月曜 9:00 ローンチMTG前まで）",""),
("","各LDシートの『ローンチ進捗』を更新 → MTGはダッシュボードの🔴→🟡→🟢の順に確認"),
("■ 自動アラート（メール）",""),
("","毎日決まった時刻にメールが届きます：①ローンチまとめ（LD＋CC宛て）②自分の担当分だけ（担当メンバー・DNA宛て）③全ローンチのまとめ（OM宛て）"),
("","アラート対象：期日超過／最終接触からN日放置／CCS後の申込待ち・振込待ち・同意書待ち／次アクション未設定／入会後タスクの期限切れ（LINEグループ未参加は初回定例会が近いと⚠付きで先頭に）"),
("","本人へのリマインドメール（LINE招待URL付き）は、行動マスタで『要』のタスクだけ。GASのCONFIGでONにします。"),
("","メールの『シートを開く』ボタンから、そのまま管理表に飛べます。アラートが0件の人にはメールは届きません。"),
("■ メールを見落とさないために（DNAにお願いすること）",""),
("","① スマホのGmailアプリで、送信元アドレスを連絡先に登録し、通知をONにする"),
("","② 件名が【要対応】【ローンチアラート】で始まるメールにスターを付けるフィルタを作る（Gmail → 設定 → フィルタ）"),
("■ ルール（RCS ローンチ方針より）",""),
("","・プリコア人数は『申込書提出・入金・プリコア同意書』の3つが揃った人のみカウント"),
("","・フェーズ1は6週間、フェーズ2は4週間が期限。期限内未達の場合は全員と面談し、①継続 ②他チャプター紹介 ③返金 の3択を提示"),
("","・フェーズ3以降はメンバー都合の返金不可。フェーズ3 Week1で必ず説明しコミットメントを取り付ける"),
("","・新メンバーには必ずウェルカムコール／1年目メンバーにはベーシック受講を確認（行動マスタに組み込み済み）"),
("■ GASの入れ方",""),
("","① Googleスプレッドシートで開く → 拡張機能 → Apps Script → Code.gs に貼り付けて保存"),
("","② シートを再読み込み → メニュー「🚀ローンチ管理」→「初期チェック」→ 権限を許可"),
("","③ 通知先シートにLDメールと、担当メンバー・DNAのメールを入力 →「アラートをプレビュー」で内容確認 →「今すぐ送信」（最初はTEST_MODEでOMにだけ届く）"),
("","④ 問題なければ Code.gs の TEST_MODE を false にして保存 →「自動実行をON」"),
("■ 新しいLD・ローンチを追加するとき",""),
("","① 既存のLDシートを複製してシート名をLD名に → ② ダッシュボードA列の空き行にLD名 → ③ 設定のLD一覧・通知先シートに追加"),
]
r = 4
for a, b in rows:
    ws.cell(r,1,a).font = font(bold=a.startswith("■"), size=10, color=NAVY if a.startswith("■") else "000000")
    c = ws.cell(r,2,b); c.font = font(size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
widths(ws, [28, 120])

# ================= 設定 =================
st = wb.create_sheet("設定")
title(st, "設定（ここを変えると全体に反映）", "A列の項目名はGASが読み取るので変更しないでください")
cfg = [("フェーズ1期限（日）",42),("フェーズ2期限（日）",28),("黄色判定ライン（目標ペース比）",0.7),
       ("期限間近アラート（残日数）",7),("ログ未入力アラート（日）",3),("CCS放置アラート（日）",3),
       ("本人リマインド再送間隔（日）",2),("初回定例会前の警戒期間（日）",7),("週次推移の開始日（月曜）",__import__("datetime").date(2026,10,5))]
hdr(st.cell(3,1,"項目")); hdr(st.cell(3,2,"値"))
for i,(k,v) in enumerate(cfg):
    cell(st,4+i,1,k); cell(st,4+i,2,v,inp=True,fmt="0%" if isinstance(v,float) else ("yyyy/mm/dd" if not isinstance(v,(int,float)) else None))
lists = [("フェーズ",["準備中","フェーズ1","フェーズ2","フェーズ3","発足済","延長判断中","解散・移籍"]),
         ("チーム",["JJ","Believe","チェキン"]),("声かけ状況",["未着手","打診中","承諾","辞退"]),
         ("候補者の状態",["進行中","保留","辞退"]),("要否",["要","不要"]),("LD一覧",names),("声かけ目標（1日）",[10]*len(names)),
         ("現状",["リストアップ済","説明会お誘い中","チャプター見学お誘い中","チャプター見学済","CCS依頼中","CCS日程確定","CCS済","検討中"]),
         ("区分",["本人","DNA"]),("担当ロール",["本人","担当","CCS実施","LD"])]
LREF = {}
for j,(name,vals) in enumerate(lists):
    col = 4+j; hdr(st.cell(3,col,name), BLUE)
    for i in range(20): cell(st,4+i,col, vals[i] if i < len(vals) else None, inp=True)
    LREF[name] = f"'設定'!${L(col)}$4:${L(col)}$23"
    st.column_dimensions[L(col)].width = 13
st.column_dimensions["A"].width = 32; st.column_dimensions["B"].width = 12
P1,P2,YLINE,NEAR,NOLOG,IDLE = ("'設定'!$B$4","'設定'!$B$5","'設定'!$B$6","'設定'!$B$7","'設定'!$B$8","'設定'!$B$9")
def dv(ws,name,rng):
    d = DataValidation(type="list", formula1="="+LREF[name], allow_blank=True); ws.add_data_validation(d); d.add(rng)
def dvcheck(ws,rng):
    d = DataValidation(type="list", formula1='"✓"', allow_blank=True); ws.add_data_validation(d); d.add(rng)

# ================= LDシート =================
ITEMS = ["管轄エリア","ローンチディレクター","チャプター名","定例会日時","初回定例会目安"] + \
        [f"CCS担当{c}" for c in "①②③④"] + [f"MS・メンター{c}" for c in "①②③④"] + \
        [f"全体サポート{c}" for c in "①②③④"] + ["その他・必要な役割"]
for name,team in LDS:
    s = wb.create_sheet(name)
    s.merge_cells("A1:E1"); s["A1"] = f"{name}｜ローンチチーム形成シート"
    s["A1"].font = font(bold=True,size=13,color="FFFFFF"); s["A1"].fill = fill(NAVY); s["A1"].alignment = Alignment(horizontal="center")
    for c,h in enumerate(["項目","内容・候補者","役割／期待すること","声かけ・進捗メモ","声かけ状況"],1): hdr(s.cell(3,c,h))
    for i,it in enumerate(ITEMS):
        r = 4+i
        cell(s,r,1,it,bold=True).fill = fill(LBLUE)
        v = team if it=="管轄エリア" else (name if it=="ローンチディレクター" else None)
        cell(s,r,2,v,inp=True,fmt="yyyy/mm/dd" if it=="初回定例会目安" else None)
        for c in (3,4): cell(s,r,c,inp=True,wrap=True)
        if r >= 9: cell(s,r,5,"未着手",inp=True,center=True)
        else: cell(s,r,5).fill = fill(GRAY)
    dv(s,"チーム","B4"); dv(s,"声かけ状況","E9:E21")
    s.merge_cells("A23:E23"); s["A23"] = "■ ローンチ進捗（毎週月曜 9:00 のローンチMTG前に更新）"; s["A23"].font = font(bold=True,color=NAVY)
    prog = [("現フェーズ","準備中",None),("フェーズ開始日",None,"yyyy/mm/dd"),("フェーズ目標人数（正式プリコア）",None,"0"),
            ("今週の課題・EDへの相談",None,None),("ローンチ詳細シートURL（NOBELIA型）",None,None)]
    for i,(k,v,f) in enumerate(prog):
        r = 24+i
        cell(s,r,1,k,bold=True).fill = fill(LBLUE)
        for c in (3,4,5): s.cell(r,c).border = BD
        s.merge_cells(start_row=r,start_column=2,end_row=r,end_column=5)
        cell(s,r,2,v,inp=True,fmt=f,wrap=True)
    dv(s,"フェーズ","B24")
    s.row_dimensions[27].height = 40
    widths(s, [32,22,34,34,12]); s.freeze_panes = "B4"
    for c,col in [("承諾","E2EFDA"),("打診中","FFF2CC"),("辞退","F8CBAD")]:
        s.conditional_formatting.add("E9:E21", FormulaRule(formula=[f'$E9="{c}"'], fill=fill(col)))

def ind(r,ref): return f"INDIRECT(\"'\"&$A{r}&\"'!{ref}\")"

# ================= CCS追っかけ =================
cs = wb.create_sheet("CCS追っかけ")
title(cs,"CCS追っかけ（リストアップ〜入会まで・全ローンチ共通・1人1行）",
      "担当＝リストアップしたメンバー。担当がCCSまで追いかけます。申込書・振込控え・プリコア同意書の3つが✓で『✅正式』（RCSローンチ方針(7)）。現状や✓を編集すると最終接触日は自動で今日（GAS）")
CH = ["No","ローンチ（LD名）","担当（リストアップしたメンバー）","見込み客名","ふりがな","カテゴリー","会社名","電話番号","メール","現状",
      "CCS日","CCS実施（DNA）","申込書","振込控え","プリコア同意書","手続き判定","正式化日","最終接触日","放置日数","次のアクション","期日","状態","アラート","行動リスト","備考"]
for c,h in enumerate(CH,1): hdr(cs.cell(4,c,h))
cs.row_dimensions[4].height = 32
N = 500; E_ = 4+N
for r in range(5,5+N):
    for c in range(1,26):
        if c == 1: cell(cs,r,c,f'=IF(D{r}="","",ROW()-4)',center=True)
        elif c == 16: cell(cs,r,c,f'=IF(D{r}="","",IF(AND(M{r}="✓",N{r}="✓",O{r}="✓"),"✅正式",IF(M{r}&N{r}&O{r}="","入会前","手続き中")))',center=True)
        elif c == 19: cell(cs,r,c,f'=IF(OR(D{r}="",R{r}=""),"",TODAY()-R{r})',center=True)
        elif c == 23:
            cell(cs,r,c,(f'=IF(OR(D{r}="",V{r}="辞退",V{r}="保留",P{r}="✅正式"),"",TRIM('
                         f'IF(AND(U{r}<>"",U{r}<TODAY()),"期日超過 ","")&IF(AND(S{r}<>"",S{r}>={IDLE}),"放置"&S{r}&"日 ","")'
                         f'&IF(AND(J{r}="CCS済",M{r}<>"✓"),"CCS後の申込待ち ","")'
                         f'&IF(AND(M{r}="✓",N{r}<>"✓"),"振込待ち ","")&IF(AND(N{r}="✓",O{r}<>"✓"),"同意書待ち ","")'
                         f'&IF(T{r}="","次アクション未設定","")))'),wrap=True)
        elif c == 24: cell(cs,r,c,center=True)
        else: cell(cs,r,c,inp=True,center=c in (10,13,14,15,22),fmt="yyyy/mm/dd" if c in (11,17,18,21) else None)
dv(cs,"LD一覧",f"B5:B{E_}"); dv(cs,"現状",f"J5:J{E_}"); dvcheck(cs,f"M5:O{E_}"); dv(cs,"候補者の状態",f"V5:V{E_}")
cs.conditional_formatting.add(f"P5:P{E_}", FormulaRule(formula=['$P5="✅正式"'], fill=fill("C6EFCE")))
cs.conditional_formatting.add(f"J5:J{E_}", FormulaRule(formula=['$J5="CCS済"'], fill=fill("DDEBF7"), font=Font(color=NAVY,bold=True)))
cs.conditional_formatting.add(f"U5:U{E_}", FormulaRule(formula=['AND($U5<>"",$U5<TODAY(),$P5<>"✅正式",$V5<>"辞退")'], fill=fill("F8CBAD"), font=Font(color="9C0006",bold=True)))
cs.conditional_formatting.add(f"W5:W{E_}", FormulaRule(formula=['$W5<>""'], fill=fill("FCE4D6"), font=Font(color="9C0006")))
cs.conditional_formatting.add(f"A5:Y{E_}", FormulaRule(formula=['$V5="辞退"'], font=Font(color="A6A6A6")))
widths(cs, [5,13,14,13,13,14,16,13,20,16,11,12,7,7,8,9,11,11,6,24,11,7,22,9,22])
cs.freeze_panes = "E5"; cs.auto_filter.ref = f"A4:Y{E_}"
CSN = "'CCS追っかけ'"

# ================= 担当別CCS（GASが作成） =================
tb = wb.create_sheet("担当別CCS")
title(tb,"担当別CCS状況（ローンチ × 担当メンバー）","メニュー「🚀ローンチ管理」→「入会者の行動リストを作成・同期」で作成／自動実行ONなら毎日自動更新。誰が何人をCCSにつないだかが分かります")

# ================= 入会後行動リスト =================
tk = wb.create_sheet("入会後行動リスト")
title(tk,"入会後行動リスト（本人のタスク＋DNAのタスク）","正式化されるとGASが行動マスタから自動作成。完了したら✓（完了日は自動）。担当者・期限は手で変えてOK")
TH = ["氏名","ローンチ（LD名）","順","タスク","区分","担当者","期限","完了","完了日","状況","メモ","キー"]
for c,h in enumerate(TH,1): hdr(tk.cell(4,c,h))
NT = 2000; TE = 4+NT
for r in range(5,5+NT):
    for c in range(1,13):
        if c == 10: cell(tk,r,c,f'=IF(D{r}="","",IF(H{r}="✓","完了",IF(G{r}="","",IF(G{r}<TODAY(),"期限超過",IF(G{r}=TODAY(),"今日まで","予定")))))',center=True)
        elif c == 12: x = cell(tk,r,c); x.font = font(size=8,color="A6A6A6")
        else: cell(tk,r,c,inp=c in (6,7,8,11),center=c in (3,5,8),fmt="yyyy/mm/dd" if c in (7,9) else None)
dvcheck(tk,f"H5:H{TE}")
for key,col,fc in [("期限超過","F8CBAD","9C0006"),("今日まで","FFEB9C","7F6000"),("完了","E7E6E6","808080")]:
    tk.conditional_formatting.add(f"J5:J{TE}", FormulaRule(formula=[f'$J5="{key}"'], fill=fill(col), font=Font(color=fc,bold=True)))
tk.conditional_formatting.add(f"E5:E{TE}", FormulaRule(formula=['$E5="DNA"'], fill=fill(LBLUE), font=Font(color=NAVY,bold=True)))
tk.conditional_formatting.add(f"A5:K{TE}", FormulaRule(formula=['$H5="✓"'], font=Font(color="A6A6A6")))
widths(tk, [13,14,5,30,7,13,11,6,11,10,26,18])
tk.freeze_panes = "B5"; tk.auto_filter.ref = f"A4:L{TE}"
TKN = "'入会後行動リスト'"

# ================= 入会後進捗（GASが作成） =================
pg = wb.create_sheet("入会後進捗")
title(pg,"入会後進捗（入会者 × 行動）","メニュー「🚀ローンチ管理」→「入会者の行動リストを作成・同期」で作成／自動実行ONなら毎日自動更新")

# ================= 行動マスタ =================
ms = wb.create_sheet("行動マスタ")
title(ms,"行動マスタ（入会後の流れ）","ハルモニアの『入会後の流れ』『メンバー表』の項目。期限＝正式化日＋日数。担当ロール：本人／担当（紹介したメンバー）／CCS実施（DNA）／LD。『本人リマインド』が要のタスクは本人へメール（GASでON時）")
MH = ["順","タスク","区分","担当ロール","期限（正式化日＋日）","本人リマインド","備考（案内・DNAの動き方）"]
for c,h in enumerate(MH,1): hdr(ms.cell(4,c,h))
MASTER = [
 (1,"ウェルカムコール","DNA","CCS実施",1,"不要","RCS：新メンバーには必ずウェルカムコール。入会後の流れを一緒に確認"),
 (2,"FB登録","本人","本人",2,"要","毎日の活動をFBに投稿。FBがない方は作成"),
 (3,"公式LINE登録","本人","本人",2,"要","友だち追加後に名前を送る"),
 (4,"FB・LINEグループ招待","DNA","LD",3,"不要","報告・班などのグループすべてに招待。参加を確認したら✓"),
 (5,"個性心理学（誕生日入力）","本人","本人",3,"要","ななメンター導入時は相性でメンターを決定"),
 (6,"リストアップ開始","本人","本人",3,"要","自分のタブに見込みになりそうな方を記入。人脈資産のデータベースづくりと伝える"),
 (7,"プロフィール写真提出","本人","本人",7,"要","ビジネスにふさわしい写真。公式LINEへ送信でOK"),
 (8,"名刺写メ提出","本人","本人",7,"要","公式LINEへ送信でOK"),
 (9,"BNIスタートガイドの視聴","本人","本人",7,"要","https://welcome.bni.jp/"),
 (10,"コネクト登録","DNA","LD",10,"不要","BNI Connectの登録を確認"),
 (11,"ショーケースチャプター参加（1回目）","本人","本人",14,"要","行きたいチャプター日時をディレクターに伝える。感想をFBでシェア"),
 (12,"略歴・GAINSシート作成","本人","本人",14,"要","1to1の前に夢・目標・人脈を記入。URLを送り合う"),
 (13,"ななメンター決定","DNA","LD",14,"不要","個性心理学の相性で決定"),
 (14,"名札制作","DNA","LD",21,"不要",""),
 (15,"ショーケースチャプター参加（2回目）","本人","本人",28,"要",""),
 (16,"メンバーサクセスオリエンテーション（MSP）受講","本人","本人",30,"要","RCS：1年目メンバーには受講を必ず確認"),
 (17,"第1回メンター","DNA","LD",30,"不要","メンターとの面談実施を確認"),
 (18,"第2回メンター","DNA","LD",60,"不要",""),
 (19,"第3回メンター","DNA","LD",90,"不要",""),
 (20,"BNI Basic受講","本人","本人",90,"要","RCS：ベーシックは最低年1回"),
 (21,"第4回メンター","DNA","LD",120,"不要",""),
 (22,"DNA面談2回目（4〜6ヶ月以内）","DNA","CCS実施",150,"不要","Give/Gainの振り返りと目標の再設定"),
 (23,"第5回メンター","DNA","LD",150,"不要",""),
 (24,"第6回メンター","DNA","LD",180,"不要",""),
 (25,"DNA面談3回目（7〜9ヶ月以内）","DNA","CCS実施",240,"不要","更新に向けた伴走。RCS：更新90日前には必ずコンタクト"),
]
for i,row in enumerate(MASTER):
    for c,v in enumerate(row,1): cell(ms,5+i,c,v,inp=True,center=c in (1,3,4,5,6),wrap=c in (2,7))
for r in range(5+len(MASTER),5+40):
    for c in range(1,8): cell(ms,r,c,inp=True,center=c in (1,3,4,5,6))
dv(ms,"区分","C5:C44"); dv(ms,"担当ロール","D5:D44"); dv(ms,"要否","F5:F44")
widths(ms, [5,38,7,11,12,11,52])

# ================= 通知先 =================
nt = wb.create_sheet("通知先")
title(nt,"通知先（アラートメールの送り先）","ローンチまとめはLDメール＋CCに届きます。右の表には担当メンバー・DNAの名前とメールを、CCS追っかけ・行動リストと同じ表記で入れてください")
NH = ["ローンチ（LD名）","LDメール","CC（複数はカンマ区切り）","メンバー用LINEグループ招待URL"]
for c,h in enumerate(NH,1): hdr(nt.cell(4,c,h))
for i in range(20):
    for c in range(1,5): cell(nt,5+i,c,names[i] if (c==1 and i<len(names)) else None,inp=True,bold=c==1)
for c,h in enumerate(["名前（担当メンバー・DNA）","メール"],6): hdr(nt.cell(4,c,h),BLUE)
for i in range(200):
    for c in range(6,8): cell(nt,5+i,c,inp=True)
widths(nt, [16,28,36,40,2,24,30])

# ================= ログ =================
al = wb.create_sheet("アラートログ")
title(al,"アラートログ（自動記録）")
for c,h in enumerate(["日時","種類","宛先","対象","内容（先頭）","キー"],1): hdr(al.cell(4,c,h))
widths(al, [17,8,24,22,60,24])

# ================= 日次ログ =================
lg = wb.create_sheet("日次ログ")
title(lg,"日次ログ（毎日1ローンチ1行・活動ゼロの日も0で入力）","数字は“その日に増えた件数”。正式プリコア数はCCS追っかけから自動集計されるので、ここには書きません")
LH = ["日付","ローンチ（LD名）","入力者","声かけ数","インフォ参加","面談実施","新規申込","新規入金","新規同意書","今日のひとこと・課題","ヘルプ要否"]
for c,h in enumerate(LH,1): hdr(lg.cell(4,c,h))
NL = 1500
for r in range(5,5+NL):
    for c in range(1,12):
        cell(lg,r,c,inp=True,fmt="yyyy/mm/dd" if c==1 else ("0" if 4<=c<=9 else None),center=c!=10)
dv(lg,"LD一覧",f"B5:B{4+NL}"); dv(lg,"要否",f"K5:K{4+NL}")
lg.conditional_formatting.add(f"A5:K{4+NL}", FormulaRule(formula=['$K5="要"'], fill=fill("FCE4D6")))
widths(lg, [11,14,12,8,9,8,8,8,9,40,8]); lg.freeze_panes = "C5"; lg.auto_filter.ref = f"A4:K{4+NL}"
LOG = "'日次ログ'"
LA, LB, LD_, LF = f"{LOG}!$A$5:$A$1504", f"{LOG}!$B$5:$B$1504", f"{LOG}!$D$5:$D$1504", f"{LOG}!$F$5:$F$1504"
CB, CC, CN, CO, CT, CJ = (f"{CSN}!${x}$5:${x}${E_}" for x in "BDPQVJ")
TB, TD, TH_, TJ = (f"{TKN}!${x}$5:${x}${TE}" for x in "BDHJ")

# ================= ダッシュボード =================
db = wb.create_sheet("📊ダッシュボード")
title(db,"📊 ローンチ ダッシュボード（全チーム）","A列にLD名（シート名と同じ）を入れるだけで自動集計。🔴→🟡→🟢の順に週次MTGで確認")
D0 = 8; DN = 20
rng = lambda col: f"{col}{D0}:{col}{D0+DN-1}"
S = ["稼働中ローンチ","正式プリコア合計","🔴 要対策","🟡 要注意","直近7日 面談合計","本日ログ入力数","ログ停滞ローンチ","入会前アラート","入会後タスク遅れ","LINEグループ未招待"]
SF = [f'=COUNTIF({rng("E")},"フェーズ*")', f'=SUM({rng("J")})', f'=COUNTIF({rng("M")},"🔴*")', f'=COUNTIF({rng("M")},"🟡*")',
      f'=SUM({rng("O")})', f'=COUNTIF({LA},TODAY())',
      f'=SUMPRODUCT((LEFT({rng("E")},4)="フェーズ")*((({rng("P")}="")+(IFERROR(TODAY()-{rng("P")},0)>={NOLOG}))>0))',
      f'=COUNTIF({CSN}!$W$5:$W${E_},"?*")', f'=COUNTIF({TJ},"期限超過")', f'=COUNTIFS({TD},"*LINEグループ*",{TH_},"<>✓")']
for i,(k,f) in enumerate(zip(S,SF)):
    c = 1+i; hdr(db.cell(4,c,k),BLUE)
    x = db.cell(5,c,f); x.font = font(bold=True,size=16,color=NAVY); x.alignment = Alignment(horizontal="center"); x.border = BD
db.row_dimensions[4].height = 30
DH = ["ローンチディレクター","チーム","チャプター名","初回定例会目安","現フェーズ","フェーズ開始日","フェーズ期限","残日数","目標人数",
      "正式プリコア","手続き中・入会前","達成率","判定","直近7日\n声かけ","直近7日\n面談","最終ログ日","体制\n(承諾/記入)","今週の課題・EDへの相談","詳細シート","入会後\nタスク遅れ","CCS済\n（累計）"]
for c,h in enumerate(DH,1): hdr(db.cell(7,c,h))
db.row_dimensions[7].height = 32
for i in range(DN):
    r = D0+i
    cell(db,r,1,names[i] if i < len(names) else None,inp=True,bold=True)
    E = lambda ref: f'IFERROR(IF({ind(r,ref)}="","",{ind(r,ref)}),"")'
    fm = {
     2:f'=IF($A{r}="","",{E("B4")})', 3:f'=IF($A{r}="","",{E("B6")})', 4:f'=IF($A{r}="","",{E("B8")})',
     5:f'=IF($A{r}="","",{E("B24")})', 6:f'=IF($A{r}="","",{E("B25")})',
     7:f'=IF(OR($A{r}="",F{r}=""),"",IF(E{r}="フェーズ1",F{r}+{P1},IF(E{r}="フェーズ2",F{r}+{P2},"")))',
     8:f'=IF(G{r}="","",G{r}-TODAY())', 9:f'=IF($A{r}="","",{E("B26")})',
     10:f'=IF($A{r}="","",COUNTIFS({CB},$A{r},{CN},"✅正式"))',
     11:f'=IF($A{r}="","",COUNTIFS({CB},$A{r},{CC},"<>",{CT},"<>辞退")-J{r})',
     12:f'=IF(OR($A{r}="",I{r}="",I{r}=0),"",J{r}/I{r})',
     13:(f'=IF(OR($A{r}="",I{r}="",F{r}=""),"",IF(J{r}>=I{r},"🟢達成",IF(G{r}="","🟡進行中",IF(H{r}<0,"🔴期限超過",'
         f'IF(AND(H{r}<={NEAR},J{r}<I{r}*0.8),"🔴期限間近",IF(J{r}>=I{r}*MIN(1,(TODAY()-F{r})/(G{r}-F{r})),"🟢順調",'
         f'IF(J{r}>=I{r}*MIN(1,(TODAY()-F{r})/(G{r}-F{r}))*{YLINE},"🟡要注意","🔴遅れ")))))))'),
     14:f'=IF($A{r}="","",SUMIFS({LD_},{LB},$A{r},{LA},">="&(TODAY()-6)))',
     15:f'=IF($A{r}="","",SUMIFS({LF},{LB},$A{r},{LA},">="&(TODAY()-6)))',
     16:f'=IF($A{r}="","",IF(COUNTIF({LB},$A{r})=0,"",_xlfn.MAXIFS({LA},{LB},$A{r})))',
     17:f'=IF($A{r}="","",IFERROR(COUNTIF({ind(r,"E9:E21")},"承諾")&" / "&COUNTA({ind(r,"B9:B21")}),""))',
     18:f'=IF($A{r}="","",{E("B27")})', 19:f'=IF($A{r}="","",{E("B28")})',
     20:f'=IF($A{r}="","",COUNTIFS({TB},$A{r},{TJ},"期限超過"))',
     21:f'=IF($A{r}="","",COUNTIFS({CB},$A{r},{CJ},"CCS済")+COUNTIFS({CB},$A{r},{CN},"✅正式",{CJ},"<>CCS済"))',
    }
    for c,f in fm.items():
        cell(db,r,c,f,center=c not in (3,18,19),wrap=c==18,fmt="yyyy/mm/dd" if c in (4,6,7,16) else ("0%" if c==12 else None))
widths(db, [14,9,16,12,10,12,12,7,7,9,10,8,12,8,8,11,10,34,14,9,9])
for key,col in [("🔴","F8CBAD"),("🟡","FFEB9C"),("🟢","C6EFCE")]:
    db.conditional_formatting.add(rng("M"), FormulaRule(formula=[f'LEFT($M{D0},1)="{key}"'], fill=fill(col)))
db.conditional_formatting.add(rng("H"), FormulaRule(formula=[f'AND(ISNUMBER($H{D0}),$H{D0}<={NEAR})'], font=Font(color="C00000",bold=True)))
db.conditional_formatting.add(rng("P"), FormulaRule(formula=[f'AND(LEFT($E{D0},4)="フェーズ",OR($P{D0}="",AND(ISNUMBER($P{D0}),TODAY()-$P{D0}>={NOLOG})))'], fill=fill("FCE4D6")))
db.conditional_formatting.add(rng("T"), FormulaRule(formula=[f'AND(ISNUMBER($T{D0}),$T{D0}>0)'], fill=fill("F8CBAD"), font=Font(color="9C0006",bold=True)))
db.freeze_panes = "B8"

# ================= 全体管理 =================
ov = wb.create_sheet("全体管理")
title(ov,"全体管理（各LDシートから自動集約・入力不要）","名前の後ろの（打診中）等は声かけ状況。承諾済みは名前のみ表示")
OH = ["ローンチディレクター","管轄エリア","チャプター名","定例会日時","初回定例会","CCS担当①","CCS担当②","CCS担当③","CCS担当④",
      "MS・メンター①","MS・メンター②","MS・メンター③","MS・メンター④","全体サポート","その他・必要な役割","進捗・メモ"]
for c,h in enumerate(OH,1): hdr(ov.cell(4,c,h))
for i in range(DN):
    r = 5+i
    cell(ov,r,1,f"='📊ダッシュボード'!A{D0+i}&\"\"",bold=True)
    def pick(row):
        b = ind(r,"B"+str(row)); e = ind(r,"E"+str(row))
        return f'=IF($A{r}="","",IFERROR(IF({b}="","",{b}&IF({e}="承諾",""," （"&{e}&"）")),""))'
    fl = lambda ref: f'=IF($A{r}="","",IFERROR(IF({ind(r,ref)}="","",{ind(r,ref)}),""))'
    vals = {2:fl("B4"),3:fl("B6"),4:fl("B7"),5:fl("B8")}
    for k,row in enumerate(range(9,17)): vals[6+k] = pick(row)
    vals[14] = f'=IF($A{r}="","",IFERROR(_xlfn.TEXTJOIN("、",TRUE,{ind(r,"B17:B20")}),""))'
    vals[15] = pick(21); vals[16] = fl("B27")
    for c,f in vals.items(): cell(ov,r,c,f,wrap=True,fmt="yyyy/mm/dd" if c==5 else None)
widths(ov, [14,10,14,14,11]+[13]*8+[22,14,30]); ov.freeze_panes = "B5"

# ================= 週次推移 =================
wk = wb.create_sheet("週次推移")
WSTART = "'設定'!$B$12"; NW = 13
title(wk,"週次推移（10/5スタート・月曜はじまり・13週＝年末まで）","開始日は『設定』シートで変更できます。黄色の列＝今週／まだ来ていない週は空欄／過ぎた週で0はオレンジ")
CS_K = f"{CSN}!$K$5:$K${E_}"
blocks = [("声かけ数（日次ログ）","log","D"),("面談実施（日次ログ）","log","F"),
          ("CCS実施数（CCS追っかけのCCS日）","ccs",None),("正式化人数（CCS追っかけの正式化日）","formal",None)]
r0 = 4
for bt,kind,colL in blocks:
    wk.cell(r0,1,"■ "+bt).font = font(bold=True,color=NAVY)
    hdr(wk.cell(r0+1,1,"週"))
    hdr(wk.cell(r0+2,1,"ローンチ（LD名）"))
    for k in range(NW):
        c = 2+k
        hdr(wk.cell(r0+1,c,f"第{k+1}週"))
        x = hdr(wk.cell(r0+2,c,f"={WSTART}+7*{k}")); x.number_format = "m/d〜"
    hdr(wk.cell(r0+1,2+NW,"")); hdr(wk.cell(r0+2,2+NW,"累計"))
    hr = r0+2
    for i in range(DN):
        r = r0+3+i
        cell(wk,r,1,f"='📊ダッシュボード'!A{D0+i}&\"\"",bold=True)
        for k in range(NW):
            h = f"{L(2+k)}${hr}"
            if kind == "log":
                core = f'SUMIFS({LOG}!${colL}$5:${colL}$1504,{LB},$A{r},{LA},">="&{h},{LA},"<"&({h}+7))'
            elif kind == "ccs":
                core = f'COUNTIFS({CB},$A{r},{CS_K},">="&{h},{CS_K},"<"&({h}+7))'
            else:
                core = f'COUNTIFS({CB},$A{r},{CN},"✅正式",{CO},">="&{h},{CO},"<"&({h}+7))'
            cell(wk,r,2+k,f'=IF(OR($A{r}="",{h}>TODAY()),"",{core})',center=True)
        cell(wk,r,2+NW,f'=IF($A{r}="","",SUM(B{r}:{L(1+NW)}{r}))',center=True,bold=True)
    tr = r0+3+DN
    cell(wk,tr,1,"リージョン合計",bold=True).fill = fill(LBLUE)
    for k in range(NW+1):
        c = 2+k
        x = cell(wk,tr,c,f'=IF({L(c)}${hr}>TODAY(),"",SUM({L(c)}{r0+3}:{L(c)}{tr-1}))' if k < NW else f'=SUM({L(c)}{r0+3}:{L(c)}{tr-1})',center=True,bold=True)
        x.fill = fill(LBLUE)
    rng_ = f"B{r0+3}:{L(1+NW)}{tr}"
    wk.conditional_formatting.add(rng_, FormulaRule(formula=[f'AND(TODAY()>=B${hr},TODAY()<B${hr}+7)'], fill=fill("FFF2CC")))
    wk.conditional_formatting.add(f"B{r0+3}:{L(1+NW)}{tr-1}", FormulaRule(formula=[f'AND($A{r0+3}<>"",ISNUMBER(B{r0+3}),B{r0+3}=0,B${hr}+7<=TODAY())'], fill=fill("FCE4D6")))
    wk.conditional_formatting.add(f"B{hr}:{L(1+NW)}{hr}", FormulaRule(formula=[f'AND(TODAY()>=B${hr},TODAY()<B${hr}+7)'], fill=fill("C8102E")))
    r0 = tr + 3
wk.column_dimensions["A"].width = 16
for c in range(2,3+NW): wk.column_dimensions[L(c)].width = 8
wk.freeze_panes = "B4"

order = ["使い方","📊ダッシュボード","全体管理","週次推移","CCS追っかけ","担当別CCS","入会後行動リスト","入会後進捗","日次ログ"]+names+ \
        ["行動マスタ","通知先","設定","アラートログ"]
wb._sheets = [wb[n] for n in order]
for s in wb.worksheets: s.sheet_view.showGridLines = False
wb.active = 1
wb.save(OUT)
print("ok")
